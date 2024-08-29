import docx

def clean_and_format_questions(file_path):
    questions = []
    doc = docx.Document(file_path)

    question_text = ""
    choices = []
    answer = ""
    explanation = ""

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Question"):
            if question_text and len(choices) == 4:  # Ensure exactly 4 choices
                questions.append((question_text, choices, answer, explanation))
            question_text = text
            choices = []
            answer = ""
            explanation = ""
        elif text.startswith("A)") or text.startswith("B)") or text.startswith("C)") or text.startswith("D)"):
            choices.append(text)  # Keep the full text for choices
        elif text.startswith("Answer:"):
            answer = text.split(":")[1].strip()
        elif text.startswith("Explanation:"):
            explanation = text.split(":", 1)[1].strip()

    if question_text and len(choices) == 4:  # Check last question
        questions.append((question_text, choices, answer, explanation))

    # Format each question for Streamlit
    formatted_questions = []
    for question_text, choices, answer, explanation in questions:
        formatted_choices = [
            f"A) {choices[0].split(') ')[1]}",
            f"B) {choices[1].split(') ')[1]}",
            f"C) {choices[2].split(') ')[1]}",
            f"D) {choices[3].split(') ')[1]}"
        ]
        formatted_questions.append((question_text, formatted_choices, answer, explanation))

    return formatted_questions

# Example usage
file_path = "mcat.docx"
formatted_questions = clean_and_format_questions(file_path)

# Display formatted questions for verification
for q in formatted_questions:
    print(q[0])  # Question text
    for choice in q[1]:  # Choices
        print(choice)
    print(f"Answer: {q[2]}")
    print(f"Explanation: {q[3]}\n")

import docx

def clean_and_format_docx(input_file, output_file):
    doc = docx.Document(input_file)
    formatted_questions = []

    question_text = ""
    choices = []
    answer = ""
    explanation = ""

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Question"):
            if question_text and len(choices) == 4:  # Ensure exactly 4 choices
                formatted_questions.append((question_text, choices, answer, explanation))
            question_text = text
            choices = []
            answer = ""
            explanation = ""
        elif text.startswith("A)") or text.startswith("B)") or text.startswith("C)") or text.startswith("D)"):
            choices.append(text)
        elif text.startswith("Answer:"):
            answer = text.split(":")[1].strip()
        elif text.startswith("Explanation:"):
            explanation = text.split(":", 1)[1].strip()

    # Check the last question
    if question_text and len(choices) == 4:
        formatted_questions.append((question_text, choices, answer, explanation))

    # Write to a new document
    new_doc = docx.Document()
    for question, choices, answer, explanation in formatted_questions:
        new_doc.add_paragraph(question)
        for choice in choices:
            new_doc.add_paragraph(choice)
        new_doc.add_paragraph(f"Answer: {answer}")
        new_doc.add_paragraph(f"Explanation: {explanation}")
        new_doc.add_paragraph()  # Add a blank line for spacing

    new_doc.save(output_file)

if __name__ == "__main__":
    input_file = "mcat.docx"  # Original file
    output_file = "formatted_mcat.docx"  # Cleaned output file
    clean_and_format_docx(input_file, output_file)
